"""Portable Excel export for the installed Python application."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from coursepulse.io import frames, safe
from coursepulse.core import MOSCOW
import pandas as pd

LABELS={'student_id':'ID слушателя','name':'Слушатель','group':'Группа','active':'Обучается',
        'expected':'Занятий прошло','attended':'Участвовал','absent':'Пропуски','excused':'Уважительные',
        'unknown':'Нет отметки','attendance_rate':'Посещаемость','due':'Работ к сроку','accepted':'Принято',
        'overdue':'Просрочено','revision':'На доработке','waiting_review':'Ждёт проверки',
        'completion_rate':'Доля принятых','missing_reasons':'Нет причины','priority':'Приоритет',
        'contact':'Контакт','reason':'Причина','action':'Действие куратора','draft':'Черновик сообщения',
        'session_id':'ID занятия','title':'Тема','start_at':'Начало (МСК)',
        'status':'Статус','assignment_id':'ID задания','due_at':'Дедлайн (МСК)',
        'state':'Состояние','submitted_at':'Сдано (МСК)','late':'Поздняя сдача','comment':'Комментарий'}


def write_excel(report, path):
    workbook = Workbook(); workbook.remove(workbook.active)
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    for name, frame in frames(report).items():
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.append([name]); sheet.append(['Состояние на', report.as_of]); sheet.append(['Снимок расчёта Python. Для обновления повторите построение отчёта.'])
        sheet.append([LABELS.get(column,column) for column in frame.columns])
        for record in frame.itertuples(index=False, name=None):
            values=[]
            for column,value in zip(frame.columns,record):
                if pd.isna(value):values.append(None)
                elif column in ['start_at','due_at','submitted_at']:
                    values.append(pd.Timestamp(value).tz_convert(MOSCOW).tz_localize(None).to_pydatetime() if value else None)
                else:values.append(safe(value))
            sheet.append(values)
        sheet.freeze_panes = 'D5'
        sheet.auto_filter.ref = f'A4:{get_column_letter(len(frame.columns))}{max(4, sheet.max_row)}'
        for cell in sheet[1]: cell.font = Font(name='Arial',size=15,bold=True,color='233E60')
        for cell in sheet[4]:
            cell.fill = PatternFill('solid',fgColor='233E60');cell.font = Font(name='Arial',color='FFFFFF',bold=True)
            cell.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
        sheet.row_dimensions[4].height = 34
        for row in sheet.iter_rows(min_row=5):
            for cell in row:
                cell.font = Font(name='Arial',size=10)
                cell.alignment = Alignment(vertical='center',wrap_text=True)
        for index, column in enumerate(frame.columns,1):
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = 32 if column in ['name','title','reason','action','comment','draft'] else 19
            if column in ['start_at','due_at','submitted_at']: sheet.column_dimensions[letter].width=30
            if column in ['start_at','due_at','submitted_at']:
                for row in range(5,sheet.max_row+1):sheet.cell(row,index).number_format='dd.mm.yyyy hh:mm'
            if column.endswith('_rate'):
                for row in range(5,sheet.max_row+1): sheet.cell(row,index).number_format='0.0%'
        if name=='Сводка':
            # Two rates stay auditable and editable in Excel; their underlying counts are snapshot inputs.
            for row in range(5,sheet.max_row+1):
                sheet.cell(row,10,f'=IF(OR(I{row}>0,E{row}-H{row}=0),"нет данных",F{row}/(E{row}-H{row}))')
                sheet.cell(row,16,f'=IF(K{row}=0,"нет данных",L{row}/K{row})')
            if sheet.max_row>=5:
                sheet.conditional_formatting.add(f'M5:M{sheet.max_row}',CellIsRule(operator='greaterThan',formula=['0'],fill=PatternFill('solid',fgColor='FCE4D6')))
    workbook.save(path)
