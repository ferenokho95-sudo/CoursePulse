import json
import tempfile
import unittest
from pathlib import Path
import pandas as pd
from coursepulse.core import Data, build, instant
from coursepulse.sample import initialize
from coursepulse.io import export, update

AS_OF='2026-10-29T18:00:00+03:00'


class TrackerTests(unittest.TestCase):
    def setUp(self):
        self.temporary=tempfile.TemporaryDirectory();self.addCleanup(self.temporary.cleanup)
        self.path=Path(self.temporary.name)/'data';initialize(self.path,demo=True)

    def report(self,point=AS_OF):return build(Data.load(self.path),point)

    def table(self,name):return pd.read_csv(self.path/f'{name}.csv',dtype=str,keep_default_na=False)

    def save(self,name,frame):frame.to_csv(self.path/f'{name}.csv',index=False,encoding='utf-8-sig')

    def test_demo_totals_and_student_metrics(self):
        report=self.report();summary=report.summary.set_index('student_id')
        self.assertEqual(len(summary),6);self.assertEqual(summary.expected.sum(),24)
        self.assertEqual(summary.loc['S002','absent'],2);self.assertEqual(summary.loc['S002','attendance_rate'],.5)
        self.assertEqual(summary.loc['S002','overdue'],2)
        self.assertEqual(summary.loc['S003','excused'],1);self.assertEqual(summary.loc['S003','attendance_rate'],1)
        self.assertEqual(summary.loc['S003','revision'],1)

    def test_unknown_is_not_absence(self):
        row=self.report().summary.set_index('student_id').loc['S004']
        self.assertEqual(row['unknown'],1);self.assertEqual(row['absent'],0)
        self.assertTrue(pd.isna(row['attendance_rate']))

    def test_late_attendance_counts_as_attended(self):
        row=self.report().summary.set_index('student_id').loc['S005']
        self.assertEqual(row['attended'],4);self.assertEqual(row['attendance_rate'],1)

    def test_submitted_work_is_not_overdue(self):
        row=self.report().summary.set_index('student_id').loc['S004']
        self.assertEqual(row['waiting_review'],1);self.assertEqual(row['overdue'],0)

    def test_future_sessions_not_in_denominator(self):
        r=self.report();self.assertFalse(r.attendance.session_id.eq('L05').any())

    def test_future_submission_does_not_change_earlier_slice(self):
        earlier=self.report('2026-10-29T09:00:00+03:00')
        state=earlier.homework.query("student_id=='S001' and assignment_id=='HW03'").iloc[0]['state']
        self.assertEqual(state,'not_started')

    def test_deadline_equality_is_still_allowed(self):
        r=self.report('2026-10-15T23:59:00+03:00')
        task=r.homework.query("student_id=='S002' and assignment_id=='HW01'").iloc[0]
        self.assertEqual(task['state'],'not_started')

    def test_enrollment_limits_expected_sessions(self):
        frame=self.table('students');frame.loc[0,'enrolled_at']='2026-10-20T00:00:00+03:00'
        self.save('students',frame)
        for name,key,allowed in [('attendance','session_id',['L04']),('submissions','assignment_id',['HW03'])]:
            frame=self.table(name);frame=frame[~frame.student_id.eq('S001') | frame[key].isin(allowed)];self.save(name,frame)
        row=self.report().summary.set_index('student_id').loc['S001']
        self.assertEqual(row['expected'],1);self.assertEqual(row['due'],0)

    def test_ended_students_not_contacted(self):
        frame=self.table('students');frame.loc[1,'ended_at']='2026-10-29T00:00:00+03:00';self.save('students',frame)
        self.assertNotIn('S002',self.report().contacts.student_id.tolist())

    def test_duplicate_keys_rejected(self):
        f=self.table('attendance');self.save('attendance',pd.concat([f,f.iloc[:1]]))
        with self.assertRaises(ValueError):Data.load(self.path)

    def test_unknown_student_rejected(self):
        f=self.table('attendance');f.loc[0,'student_id']='missing';self.save('attendance',f)
        with self.assertRaises(ValueError):Data.load(self.path)

    def test_invalid_status_rejected(self):
        f=self.table('attendance');f.loc[0,'status']='whatever';self.save('attendance',f)
        with self.assertRaises(ValueError):Data.load(self.path)

    def test_naive_dates_rejected(self):
        with self.assertRaises(ValueError):instant('2026-10-29')

    def test_invalid_deadline_rejected(self):
        f=self.table('assignments');f.loc[0,'due_at']='2026-01-01T00:00:00+03:00';self.save('assignments',f)
        with self.assertRaises(ValueError):Data.load(self.path)

    def test_contact_priority_and_drafts(self):
        r=self.report();self.assertEqual(r.contacts.iloc[0].student_id,'S002')
        self.assertEqual(r.contacts.iloc[0].priority,'высокий')
        self.assertEqual(len(r.contacts),5);self.assertIn('Уточнить',r.contacts.iloc[0].action)

    def test_atomic_update_changes_existing_mark(self):
        original=len(self.table('attendance'))
        update(self.path,'attendance',{'student_id':'S002','session_id':'L02','status':'present','reason':''})
        self.assertEqual(len(self.table('attendance')),original)
        row=self.report().summary.set_index('student_id').loc['S002'];self.assertEqual(row['absent'],1)

    def test_bad_update_does_not_change_file(self):
        before=(self.path/'attendance.csv').read_bytes()
        with self.assertRaises(ValueError):update(self.path,'attendance',{'student_id':'missing','session_id':'L02','status':'present','reason':''})
        self.assertEqual(before,(self.path/'attendance.csv').read_bytes())

    def test_empty_template_is_valid(self):
        target=Path(self.temporary.name)/'empty';initialize(target)
        r=build(Data.load(target),AS_OF);self.assertTrue(r.summary.empty)

    def test_demo_init_never_overwrites(self):
        with self.assertRaises(ValueError):initialize(self.path,demo=True)

    def test_export_html_escaping_and_csv_formula_safety(self):
        f=self.table('students');f.loc[1,'name']='<script>alert(1)</script>';f.loc[1,'contact']='=HYPERLINK("bad")';self.save('students',f)
        out=Path(self.temporary.name)/'reports';export(self.report(),out,excel=False)
        self.assertNotIn('<script>',(out/'report.html').read_text())
        self.assertIn('&lt;script&gt;', (out/'report.html').read_text())
        exported=pd.read_csv(out/'contacts.csv',keep_default_na=False)
        self.assertTrue(exported.query("student_id=='S002'").iloc[0].contact.startswith("'="))
        self.assertEqual(len(json.loads((out/'report.json').read_text())['Сводка']),6)

    def test_portable_excel_export(self):
        from openpyxl import load_workbook
        out=Path(self.temporary.name)/'xlsx'
        export(self.report(),out,excel=True)
        workbook=load_workbook(out/'course_report.xlsx',data_only=False)
        self.assertEqual(workbook.sheetnames,['Сводка','Связаться','Посещаемость','Домашние работы'])
        sheet=workbook['Сводка']
        self.assertEqual(sheet['A4'].value,'ID слушателя')
        self.assertEqual(sheet['G6'].value,2)
        self.assertEqual(sheet['M6'].value,2)
        self.assertTrue(sheet['J6'].value.startswith('=IF('))
        self.assertEqual(workbook['Посещаемость']['E5'].value.hour,19)
        self.assertIsNotNone(sheet.auto_filter.ref)
        self.assertTrue(workbook.calculation.fullCalcOnLoad)


if __name__=='__main__':unittest.main()
